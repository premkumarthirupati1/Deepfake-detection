from io import BytesIO
from django.shortcuts import render, redirect
from django.contrib import messages
from django.db.models import Count, Avg, Q
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from Remote_User.models import ClientRegister_Model, deepfake_detection_type, detection_ratio, detection_accuracy
from deepfake_detection.ml_engine import train_and_evaluate_all, load_trained_models


def serviceproviderlogin(request):
    """Admin / Service Provider Login."""
    if request.method == "POST":
        admin = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()

        if admin == "Admin" and password == "Admin":
            request.session['admin_logged_in'] = True
            request.session['admin_name'] = "Admin"
            messages.success(request, "Welcome, Administrator!")
            return redirect('View_Remote_Users')
        else:
            messages.error(request, "Invalid Admin Credentials.")

    return render(request, 'SProvider/serviceproviderlogin.html')


def serviceproviderlogout(request):
    """Admin Logout."""
    if 'admin_logged_in' in request.session:
        del request.session['admin_logged_in']
    messages.info(request, "Admin logged out successfully.")
    return redirect('login')


def View_Deepfake_Prediction_Ratio(request):
    """Computes and renders the current Real vs Fake prediction ratio."""
    total_count = deepfake_detection_type.objects.count()
    detection_ratio.objects.all().delete()

    fake_count = deepfake_detection_type.objects.filter(Prediction__iexact='FAKE').count()
    real_count = deepfake_detection_type.objects.filter(Prediction__iexact='REAL').count()

    if total_count > 0:
        fake_ratio = round((fake_count / total_count) * 100, 2)
        real_ratio = round((real_count / total_count) * 100, 2)

        detection_ratio.objects.create(names='FAKE', ratio=str(fake_ratio))
        detection_ratio.objects.create(names='REAL', ratio=str(real_ratio))
    else:
        fake_ratio = 0.0
        real_ratio = 0.0

    objs = detection_ratio.objects.all()
    context = {
        'objs': objs,
        'total_count': total_count,
        'fake_count': fake_count,
        'real_count': real_count,
        'fake_ratio': fake_ratio,
        'real_ratio': real_ratio
    }
    return render(request, 'SProvider/View_Deepfake_Prediction_Ratio.html', context)


def View_Remote_Users(request):
    """Lists all registered users."""
    users = ClientRegister_Model.objects.all().order_by('-id')
    return render(request, 'SProvider/View_Remote_Users.html', {'objects': users})


def View_deepfake_prediction_Type(request):
    """Lists all prediction records with status badges."""
    records = deepfake_detection_type.objects.all().order_by('-id')
    return render(request, 'SProvider/View_deepfake_prediction_Type.html', {'list_objects': records})


def train_model(request):
    """
    Executes training of all forensic ML classifiers, updates database accuracy records,
    and returns comprehensive metrics.
    """
    metrics = train_and_evaluate_all()

    detection_accuracy.objects.all().delete()
    for name, data in metrics.items():
        detection_accuracy.objects.create(names=name, ratio=str(data['accuracy']))

    objs = detection_accuracy.objects.all()
    messages.success(request, "All Machine Learning Models Successfully Retrained & Evaluated!")
    return render(request, 'SProvider/train_model.html', {'objs': objs, 'metrics': metrics})


def charts(request, chart_type='pie'):
    """Renders Real vs. Fake prediction ratio charts."""
    # Ensure ratio data exists
    if not detection_ratio.objects.exists():
        View_Deepfake_Prediction_Ratio(request)

    ratios = detection_ratio.objects.all()
    chart_data = [{'names': r.names, 'dcount': float(r.ratio)} for r in ratios]

    return render(request, 'SProvider/charts.html', {
        'form': chart_data,
        'chart_type': chart_type
    })


def charts1(request, chart_type='bar'):
    """Renders Model accuracy comparison charts."""
    if not detection_accuracy.objects.exists():
        # Prepopulate or train
        train_model(request)

    accuracies = detection_accuracy.objects.all()
    chart_data = [{'names': a.names, 'dcount': float(a.ratio)} for a in accuracies]

    return render(request, 'SProvider/charts1.html', {
        'form': chart_data,
        'chart_type': chart_type
    })


def likeschart(request, like_chart='bar'):
    """Alternate bar chart view for model accuracies."""
    return charts1(request, chart_type=like_chart)


def Download_Trained_DataSets(request):
    """
    Exports prediction records into a beautifully styled modern Excel (.xlsx) file.
    """
    records = deepfake_detection_type.objects.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Deepfake Predictions"

    # Styling definitions
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    headers = [
        "ID", "URL / Source", "IP Address", "Video Filename",
        "Width (px)", "Height (px)", "Country", "Locale",
        "Latitude", "Longitude", "Prediction"
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Add data rows
    for r in records:
        row_data = [
            r.id, r.url, r.ipaddress, r.videoname,
            r.original_width, r.original_height, r.country, r.locale,
            r.latitude, r.longitude, r.Prediction
        ]
        ws.append(row_data)

    # Style rows and set auto column widths
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            if cell.column == 11:  # Prediction
                cell.alignment = center_align
                if str(cell.value).upper() == 'REAL':
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="047857")
                else:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Deepfake_Detection_Predictions.xlsx"'
    return response